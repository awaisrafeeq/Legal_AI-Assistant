"""
OCR script for internal documents container.
Handles ALL file types: PDF, Images, Excel, CSV, Word, Text, Audio/Video.
Reads from INTERNAL_CONTAINER_SAS_URL, writes OCR output to CONTAINER_SAS_URL
under extracted-text/internal/ with Source: header for correct metadata extraction.
"""
import os
import sys
import logging
from typing import Optional, Tuple
from urllib.parse import quote, urlparse
from pathlib import Path

from azure.ai.formrecognizer import DocumentAnalysisClient
from azure.core.credentials import AzureKeyCredential
from azure.core.exceptions import HttpResponseError, ResourceNotFoundError
from azure.storage.blob import BlobServiceClient
from dotenv import load_dotenv
from tenacity import retry, stop_after_attempt, wait_exponential, retry_if_exception_type
from tqdm import tqdm

sys.stdout.reconfigure(encoding='utf-8')

os.makedirs("logs", exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler("logs/ocr_internal.log", encoding="utf-8"),
    ]
)
logger = logging.getLogger(__name__)

load_dotenv()

# ── File type categories ──────────────────────────────────────────────
AUDIO_EXTS = {'.mp3', '.wav', '.m4a', '.flac', '.aac', '.ogg', '.wma'}
VIDEO_EXTS = {'.mp4', '.avi', '.mov', '.wmv', '.mkv', '.flv', '.webm'}
EXCEL_EXTS = {'.xlsx', '.xls', '.xlsm'}
WORD_EXTS = {'.docx', '.doc'}
IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.tiff', '.tif', '.bmp', '.gif', '.webp'}
PDF_EXTS = {'.pdf'}
CSV_EXTS = {'.csv'}
TEXT_EXTS = {'.txt', '.htm', '.html', '.xml', '.json', '.md'}
SKIP_EXTS = {'.ds_store', '.db', '.ini', '.log'}


def split_container_sas_url(container_sas_url: str, fallback_container: Optional[str] = None) -> Tuple[str, str, str]:
    p = urlparse(container_sas_url)
    if not p.scheme or not p.netloc:
        raise ValueError("URL must be a full https URL")
    if not p.query:
        raise ValueError("URL must include SAS query string")
    container_name = p.path.strip("/").split("/")[-1] if p.path.strip("/") else (fallback_container or "")
    if not container_name:
        raise ValueError("URL path must end with container name or OUTPUT_CONTAINER must be set")
    account_url = f"{p.scheme}://{p.netloc}"
    return account_url, container_name, p.query


def blob_url(account_url: str, container: str, sas: str, blob_name: str) -> str:
    safe = quote(blob_name, safe="/-_.~ ").replace(" ", "%20")
    return f"{account_url}/{container}/{safe}?{sas}"


def out_blob_name(output_prefix: str, in_blob_name: str) -> str:
    """Output goes to extracted-text/internal/{original_path}.txt"""
    op = output_prefix.strip("/")
    if op:
        return f"{op}/internal/{in_blob_name}.txt"
    return f"internal/{in_blob_name}.txt"


def get_file_category(filename: str) -> str:
    basename = Path(filename).name.lower()
    if basename in {'.ds_store', 'thumbs.db', 'desktop.ini', '.gitkeep'}:
        return 'skip'
    ext = Path(filename).suffix.lower()
    if ext in PDF_EXTS:
        return 'pdf'
    elif ext in IMAGE_EXTS:
        return 'image'
    elif ext in AUDIO_EXTS:
        return 'audio'
    elif ext in VIDEO_EXTS:
        return 'video'
    elif ext in EXCEL_EXTS:
        return 'excel'
    elif ext in CSV_EXTS:
        return 'csv'
    elif ext in WORD_EXTS:
        return 'word'
    elif ext in TEXT_EXTS:
        return 'text'
    elif ext in SKIP_EXTS:
        return 'skip'
    else:
        return 'other'


# ── Azure Document Intelligence ───────────────────────────────────────

@retry(
    wait=wait_exponential(multiplier=1, min=4, max=30),
    stop=stop_after_attempt(3),
    retry=retry_if_exception_type((HttpResponseError, ResourceNotFoundError)),
    reraise=True
)
def analyze_with_di(di_client: DocumentAnalysisClient, document_url: str):
    poller = di_client.begin_analyze_document_from_url(
        "prebuilt-read", document_url, polling_interval=5
    )
    return poller.result(timeout=300)


# ── Per-format processors ─────────────────────────────────────────────

def process_pdf_image(di_client, doc_url: str, blob_name: str) -> Optional[str]:
    try:
        result = analyze_with_di(di_client, doc_url)
        text = "\n\n".join(
            page_text for page in result.pages if (page_text := " ".join(
                line.content for line in page.lines
            ))
        )
        return text if text else f"[No text extracted from {blob_name}]"
    except Exception as e:
        logger.error(f"DI failed for {blob_name}: {e}")
        return None


def process_excel(blob_client) -> Optional[str]:
    try:
        import pandas as pd
        from io import BytesIO

        data = blob_client.download_blob().readall()
        excel_file = BytesIO(data)

        text_parts = []
        xl = pd.ExcelFile(excel_file)
        for sheet_name in xl.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            text_parts.append(f"=== Sheet: {sheet_name} ===")
            text_parts.append(df.to_string(index=False))
            text_parts.append("")

        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"Excel processing failed: {e}")
        return None


def process_csv(blob_client) -> Optional[str]:
    try:
        import pandas as pd
        from io import StringIO

        data = blob_client.download_blob().readall().decode('utf-8', errors='ignore')
        df = pd.read_csv(StringIO(data))
        return df.to_string(index=False)
    except Exception as e:
        logger.error(f"CSV processing failed: {e}")
        return None


def process_word(blob_client) -> Optional[str]:
    try:
        from docx import Document
        from io import BytesIO

        data = blob_client.download_blob().readall()
        doc = Document(BytesIO(data))

        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                text_parts.append(" | ".join(row_text))

        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"Word processing failed: {e}")
        return None


def process_text_file(blob_client) -> Optional[str]:
    try:
        data = blob_client.download_blob().readall()
        for encoding in ['utf-8', 'latin-1', 'cp1252']:
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode('utf-8', errors='ignore')
    except Exception as e:
        logger.error(f"Text file processing failed: {e}")
        return None


def process_audio_video(blob_client, blob_name: str) -> Optional[str]:
    logger.info(f"Audio/Video file detected: {blob_name}")
    size = blob_client.get_blob_properties().size
    return f"[Audio/Video file: {blob_name}. Transcription requires Azure Speech Services. Size: {size} bytes]"


def process_other(blob_client, blob_name: str) -> Optional[str]:
    try:
        props = blob_client.get_blob_properties()
        return f"[File: {blob_name}, Size: {props.size} bytes, Content-Type: {props.content_settings.content_type}]"
    except Exception as e:
        return f"[File: {blob_name} - Metadata extraction failed: {e}]"


# ── Main ──────────────────────────────────────────────────────────────

def main():
    internal_sas_url = os.environ.get("INTERNAL_CONTAINER_SAS_URL")
    if not internal_sas_url:
        logger.error("INTERNAL_CONTAINER_SAS_URL not set in .env")
        sys.exit(1)

    di_endpoint = os.environ.get("DI_ENDPOINT")
    di_key = os.environ.get("DI_KEY")
    input_prefix = os.environ.get("INPUT_PREFIX", "")
    output_prefix = os.environ.get("OUTPUT_PREFIX", "extracted-text")
    max_files_str = os.environ.get("MAX_FILES", "")
    max_files = int(max_files_str) if max_files_str.strip() else None
    overwrite = os.environ.get("OVERWRITE", "0") == "1"

    # Internal container — source files
    in_account_url, in_container, in_sas = split_container_sas_url(internal_sas_url)
    in_blob_service = BlobServiceClient(account_url=in_account_url, credential=in_sas)
    in_container_client = in_blob_service.get_container_client(in_container)

    # External container — OCR output destination
    output_sas_url = os.environ.get("CONTAINER_SAS_URL")
    output_container = os.environ.get("OUTPUT_CONTAINER") or os.environ.get("OCR_OUTPUT_CONTAINER")
    out_account_url, out_container, out_sas = split_container_sas_url(output_sas_url, output_container)
    out_blob_service = BlobServiceClient(account_url=out_account_url, credential=out_sas)

    # DI client
    di_client = DocumentAnalysisClient(
        endpoint=di_endpoint,
        credential=AzureKeyCredential(di_key)
    )

    # List all files in internal container
    prefix = input_prefix.strip("/") + "/" if input_prefix else ""
    all_blobs = []

    logger.info(f"Listing files in internal container: {in_container}")
    for blob in in_container_client.list_blobs(name_starts_with=prefix or None):
        if blob.name.startswith('.') or '/.' in blob.name:
            continue
        cat = get_file_category(blob.name)
        if cat == 'skip':
            continue
        all_blobs.append(blob.name)
        if max_files and len(all_blobs) >= max_files:
            break

    if not all_blobs:
        logger.warning("No files found in internal container")
        return

    logger.info(f"Found {len(all_blobs)} files to process")

    # Show breakdown
    by_category = {}
    for name in all_blobs:
        cat = get_file_category(name)
        by_category.setdefault(cat, []).append(name)

    logger.info("File breakdown:")
    for cat, files in sorted(by_category.items(), key=lambda x: -len(x[1])):
        logger.info(f"  {cat:10}: {len(files):5} files")

    # Process
    success_count = 0
    skip_count = 0
    fail_count = 0

    for blob_name in tqdm(all_blobs, desc="OCR Internal"):
        try:
            category = get_file_category(blob_name)
            out_name = out_blob_name(output_prefix, blob_name)
            out_client = out_blob_service.get_blob_client(out_container, out_name)

            if not overwrite:
                try:
                    out_client.get_blob_properties()
                    skip_count += 1
                    continue
                except Exception:
                    pass

            # Process based on category
            if category in ('pdf', 'image'):
                doc_url = blob_url(in_account_url, in_container, in_sas, blob_name)
                text = process_pdf_image(di_client, doc_url, blob_name)
            elif category == 'excel':
                bc = in_container_client.get_blob_client(blob_name)
                text = process_excel(bc)
            elif category == 'csv':
                bc = in_container_client.get_blob_client(blob_name)
                text = process_csv(bc)
            elif category == 'word':
                bc = in_container_client.get_blob_client(blob_name)
                text = process_word(bc)
            elif category == 'text':
                bc = in_container_client.get_blob_client(blob_name)
                text = process_text_file(bc)
            elif category in ('audio', 'video'):
                bc = in_container_client.get_blob_client(blob_name)
                text = process_audio_video(bc, blob_name)
            else:
                bc = in_container_client.get_blob_client(blob_name)
                text = process_other(bc, blob_name)

            if text:
                # Add Source header — same format as process_all_formats.py
                # This lets index_internal.py extract correct folder_path/file_name
                full_text = f"Source: {blob_name}\nType: {category.upper()}\n{'='*60}\n\n{text}"
                out_client.upload_blob(full_text.encode("utf-8"), overwrite=True)
                success_count += 1
            else:
                fail_count += 1

        except Exception as e:
            logger.error(f"Error processing {blob_name}: {e}")
            fail_count += 1
            continue

    logger.info(f"Processing complete!")
    logger.info(f"  Success: {success_count}")
    logger.info(f"  Skipped: {skip_count}")
    logger.info(f"  Failed:  {fail_count}")


if __name__ == "__main__":
    main()
